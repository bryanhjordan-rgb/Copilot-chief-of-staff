#!/usr/bin/env python3
"""Build ChiefOfStaff.xlsx from m365/schema/workbook.yaml.

Usage:
  python3 m365/scripts/build_workbook.py [output.xlsx] [--goals goals.yaml]
  python3 m365/scripts/build_workbook.py test.xlsx --fixtures m365/evals/fixtures.yaml [--today YYYY-MM-DD]

Default output: m365/workbook/ChiefOfStaff.xlsx. Upload it to OneDrive and point
the agent's Excel Online (Business) tools at it.

--goals fills the Goals table from a goals.yaml in the root kit's format:
objectives become G1, G2, ... (in priority order) and personal goals PG1, PG2, ...
Build to a path outside the repo (e.g. ~/ChiefOfStaff.xlsx) so your goals aren't
committed by accident.

--fixtures loads the eval seed data (Goals + Items) instead of the example row,
resolving TODAY, TODAY+n and TODAY-n against --today (default: the real date).
"""

import argparse
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

try:
    import yaml
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:
    sys.exit("PyYAML and openpyxl are required: pip install pyyaml openpyxl")

M365_DIR = Path(__file__).resolve().parent.parent
SCHEMA = M365_DIR / "schema" / "workbook.yaml"
DEFAULT_OUT = M365_DIR / "workbook" / "ChiefOfStaff.xlsx"

VALIDATION_ROWS = 2000
GROUP_COLORS = {
    "plan": "1F4E79",
    "why": "7F6000",
    "context": "385723",
    "tracking": "595959",
    "record": "7030A0",
    None: "1F4E79",
}
NUMERIC_TYPES = {"int", "decimal"}
FIXED_TIMESTAMP = datetime(2026, 1, 1)  # keeps rebuilt files byte-stable

# One fully worked row so the format of every multi-line cell is visible.
# Open=No keeps it out of the agent's "Open eq 'Yes'" queries.
EXAMPLE_ITEM = {
    "ID": "I-EXAMPLE",
    "Open": "No",
    "Status": "done",
    "Tier": "P2",
    "Score": 61,
    "Title": "Example: prep pricing review for Contoso renewal",
    "NextAction": "(example row - safe to delete)",
    "DueDate": "2026-11-05",
    "DoDate": "2026-11-03",
    "PlannedSlot": "09:00-10:00",
    "EffortMin": 60,
    "Energy": "deep",
    "Goal": "G1",
    "Owner": "me",
    "Rationale": "Due in 2 days, tied to the revenue goal, and the account team is waiting on it.",
    "ScoreBreakdown": "deadline=20; goal=18; stakeholder=10; unblocks=8; leverage=5; age=0",
    "Type": "meeting_prep",
    "Source": "email",
    "SourceLink": "[email] RE: Contoso renewal pricing | https://outlook.office.com/mail/...",
    "ContextLinks": (
        "[meeting] Contoso renewal review (Nov 5) | https://teams.microsoft.com/l/meetup-join/...\n"
        "[file] Contoso pricing model.xlsx | https://contoso.sharepoint.com/...\n"
        "[person] Account exec | ae@contoso.com"
    ),
    "People": "ae@contoso.com\nfinance.lead@contoso.com",
    "ResearchNotes": (
        "- Customer asked for 3-year pricing option\n"
        "- Finance approved max 12% discount (email, Oct 28)\n"
        "- Open: whether SSO tier is bundled"
    ),
    "Confidential": "No",
    "Created": "2026-10-29T08:12",
    "Updated": "2026-11-03T10:05",
    "History": (
        "2026-10-29 08:12 agent created from email\n"
        "2026-10-29 08:13 agent investigated (3 links)\n"
        "2026-11-03 10:05 agent reconciled"
    ),
    "CompletedOn": "2026-11-03T10:05",
    "Outcome": "Pricing options deck sent to account exec; 3-year option at 10% recommended.",
    "Decisions": "Bundle SSO in 3-year tier only.",
    "Artifacts": "[file] Contoso pricing options.pptx | https://contoso.sharepoint.com/...",
    "MinutesSpent": 55,
    "Steps": (
        "read pricing thread | outlook | 10\n"
        "update pricing model | excel | 25\n"
        "build options slide | powerpoint | 15\n"
        "draft cover email | outlook | ~5"
    ),
    "Signature": "prepare|pricing-review|customer",
}


def goal_rows(goals_path):
    """Map a root-kit goals.yaml onto Goals table rows."""
    data = yaml.safe_load(Path(goals_path).read_text(encoding="utf-8")) or {}
    quarter = str(data.get("quarter", ""))
    updated = str(data.get("last_updated", ""))
    objectives = sorted(data.get("objectives") or [], key=lambda o: o.get("priority", 99))
    rows = []
    for i, obj in enumerate(objectives, start=1):
        rows.append({
            "ID": f"G{i}",
            "Goal": obj.get("name", ""),
            "Priority": obj.get("priority", i),
            "Target": obj.get("target", ""),
            "KeyResults": "\n".join(obj.get("key_results") or []),
            "Progress": obj.get("progress", 0.0),
            "Status": obj.get("status", "on_track"),
            "Quarter": quarter,
            "LastUpdate": updated,
        })
    for i, goal in enumerate(data.get("personal") or [], start=1):
        rows.append({
            "ID": f"PG{i}",
            "Goal": goal.get("name", ""),
            "Progress": goal.get("progress", 0.0),
            "Status": goal.get("status", "on_track"),
            "Quarter": quarter,
            "LastUpdate": updated,
        })
    examples = [r["Goal"] for r in rows if str(r["Goal"]).lower().startswith("example")]
    if examples:
        print(f"warning: {len(examples)} goal(s) still look like template examples: {examples}")
    return rows


RELATIVE_DAY = re.compile(r"^TODAY([+-]\d+)?$")


def fixture_rows(fixtures_path, schema, today):
    """Eval fixtures with TODAY±n resolved; datetime columns get a 09:00 time."""
    data = yaml.safe_load(Path(fixtures_path).read_text(encoding="utf-8")) or {}
    seed = {}
    for table, key in (("Goals", "goals"), ("Items", "items")):
        types = {c["name"]: c.get("type") for c in schema["tables"][table]["columns"]}
        rows = []
        for row in data.get(key) or []:
            resolved = {}
            for col, value in row.items():
                if col not in types:
                    sys.exit(f"fixtures: {table} has no column '{col}'")
                match = RELATIVE_DAY.match(str(value))
                if match:
                    day = today + timedelta(days=int(match.group(1) or 0))
                    value = day.isoformat() + ("T09:00" if types[col] == "datetime" else "")
                resolved[col] = value
            rows.append(resolved)
        seed[table] = rows
    return seed


def build(out_path, goals_path=None, fixtures_path=None, today=None):
    schema = yaml.safe_load(SCHEMA.read_text(encoding="utf-8"))
    seed = {"Items": [EXAMPLE_ITEM]}
    if goals_path:
        seed["Goals"] = goal_rows(goals_path)
    if fixtures_path:
        seed = fixture_rows(fixtures_path, schema, today or date.today())
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Chief of Staff agent"
    wb.properties.created = FIXED_TIMESTAMP
    wb.properties.modified = FIXED_TIMESTAMP

    for table_name, spec in schema["tables"].items():
        add_table(wb, table_name, spec, seed.get(table_name, []))
    add_guide(wb, schema)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"wrote {out_path}")


def add_table(wb, table_name, spec, rows):
    ws = wb.create_sheet(spec["sheet"])
    columns = spec["columns"]
    names = [c["name"] for c in columns]
    if names[0] != spec["key"]:
        sys.exit(f"{table_name}: key column '{spec['key']}' must be first")
    if any(not n.isalnum() for n in names):
        sys.exit(f"{table_name}: column names must be alphanumeric: {names}")

    rows = rows or [{}]  # an Excel table needs at least one data row
    for idx, col in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        header = ws.cell(row=1, column=idx, value=col["name"])
        header.font = Font(bold=True, color="FFFFFF")
        header.fill = PatternFill("solid", fgColor=GROUP_COLORS[col.get("group")])
        header.alignment = Alignment(vertical="center")

        # Text format keeps ISO dates and IDs from being converted by Excel.
        number_format = "0.00" if col.get("type") == "decimal" else "0" if col.get("type") == "int" else "@"
        for r, row in enumerate(rows, start=2):
            cell = ws.cell(row=r, column=idx, value=row.get(col["name"]))
            cell.number_format = number_format
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[letter].width = col.get("width", 15)

        if "values" in col:
            dv = DataValidation(
                type="list",
                formula1='"' + ",".join(col["values"]) + '"',
                allow_blank=True,
            )
            dv.add(f"{letter}2:{letter}{VALIDATION_ROWS}")
            ws.add_data_validation(dv)

    last_col = get_column_letter(len(columns))
    table = Table(displayName=table_name, ref=f"A1:{last_col}{len(rows) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "G2" if table_name == "Items" else "B2"


def add_guide(wb, schema):
    ws = wb.create_sheet("Guide")
    rows = [
        ["Chief of Staff workbook", "", ""],
        [f"Schema v{schema['version']} - generated from m365/schema/workbook.yaml. Edit the schema, not this sheet.", "", ""],
        ["Rules: edit through the agent or Excel for the web; don't edit while the agent is writing.", "", ""],
        ["Rows with an empty key column are ignored by the agent. I-EXAMPLE shows the cell formats and can be deleted.", "", ""],
        ["", "", ""],
        ["Table / Column", "Allowed values", "Notes"],
    ]
    for table_name, spec in schema["tables"].items():
        rows.append([f"{table_name}  (sheet: {spec['sheet']}, key: {spec['key']})", "", spec["description"]])
        for col in spec["columns"]:
            rows.append([f"    {col['name']}", ", ".join(col.get("values", [])), col.get("note", "")])
    for row in rows:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=14)
    for cell in ws[6]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=7):
        if not row[0].value.startswith("    "):
            for cell in row:
                cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 80


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build ChiefOfStaff.xlsx from the schema.")
    parser.add_argument("output", nargs="?", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--goals", type=Path, help="goals.yaml to load into the Goals table")
    parser.add_argument("--fixtures", type=Path, help="eval fixtures.yaml to load (replaces --goals and the example row)")
    parser.add_argument("--today", type=date.fromisoformat, help="date that TODAY resolves to (default: today)")
    args = parser.parse_args()
    build(args.output, args.goals, args.fixtures, args.today)
