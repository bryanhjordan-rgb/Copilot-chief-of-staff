#!/usr/bin/env python3
"""Mine completed Items for repeated work (skill candidates).

Usage:
  python3 m365/scripts/mine.py path/to/ChiefOfStaff.xlsx --as-of 2026-10-29 [--window 30]

Reference implementation of m365/office-scripts/mine-signatures.ts, the script
the weekly flow runs inside Excel. The two must produce identical JSON; the
unittest suite runs both on the same rows. Rounding mirrors JavaScript's
Math.round (half up), not Python's round (half to even).
"""

import argparse
import json
import math
import re
import sys
from datetime import date
from pathlib import Path

JUDGMENT_TOOLS = {"judgment", "call", "meeting", "in-person", "manual"}
MAX_CANDIDATES = 3
EXAMPLE_ID = "I-EXAMPLE"
LEADING_NUMBER = re.compile(r"^\s*(-?\d+(\.\d+)?)")


def js_round(n):
    return math.floor(n + 0.5)


def round1(n):
    return js_round(n * 10) / 10


def round2(n):
    return js_round(n * 100) / 100


def day_number(iso):
    try:
        return date.fromisoformat(iso[:10]).toordinal()
    except ValueError:
        return None


def parse_float(text):
    match = LEADING_NUMBER.match(text)
    return float(match.group(1)) if match else None


def parse_steps(text):
    steps = []
    for line in text.split("\n"):
        parts = [p.strip().lower() for p in line.split("|")]
        if parts[0]:
            steps.append((parts[0], parts[1] if len(parts) > 1 else ""))
    return steps


def common_steps(group):
    """Steps (name + tool) present in at least two thirds of the group, in the most recent item's order."""
    needed = math.ceil(len(group) * 2 / 3)
    counts = {}
    for item in group:
        for step in set(item["steps"]):
            counts[step] = counts.get(step, 0) + 1
    ordered = []
    for item in group:
        for step in item["steps"]:
            if counts[step] >= needed and step not in ordered:
                ordered.append(step)
    return ordered


def detect_cadence(dates):
    if len(dates) < 2:
        return "none"
    days = sorted(day_number(d) for d in dates)
    gaps = [b - a for a, b in zip(days, days[1:])]
    if all(6 <= g <= 8 for g in gaps):
        return "weekly"
    if all(27 <= g <= 32 for g in gaps):
        return "monthly"
    return "none"


def active_signatures(skills):
    return [
        line.strip()
        for row in skills
        if row.get("Status") == "active"
        for line in row.get("Signatures", "").split("\n")
        if line.strip()
    ]


def mine(items, skills, as_of, window_days=30):
    end = day_number(as_of)
    start = end - window_days
    covered = active_signatures(skills)

    groups = {}
    considered = 0
    for row in items:
        signature = row.get("Signature", "")
        completed_on = row.get("CompletedOn", "")[:10]
        if row.get("Status") != "done" or not signature or not completed_on or row.get("ID") == EXAMPLE_ID:
            continue
        day = day_number(completed_on)
        if day is None or day < start or day > end:
            continue
        considered += 1
        groups.setdefault(signature, []).append({
            "id": row.get("ID", ""),
            "completedOn": completed_on,
            "minutes": parse_float(row.get("MinutesSpent", "")),
            "steps": parse_steps(row.get("Steps", "")),
        })

    candidates, covered_found = [], []
    for signature in sorted(groups):
        group = sorted(groups[signature], key=lambda g: g["completedOn"], reverse=True)
        if signature in covered:
            covered_found.append(signature)
            continue
        timed = [g["minutes"] for g in group if g["minutes"] is not None]
        avg = round1(sum(timed) / len(timed)) if timed else 0
        cadence = detect_cadence([g["completedOn"] for g in group])
        if not (len(group) >= 3 or (len(group) >= 2 and avg >= 20) or cadence != "none"):
            continue
        common = common_steps(group)
        tool_steps = [s for s in common if s[1] and s[1] not in JUDGMENT_TOOLS]
        fit = round2(len(tool_steps) / len(common)) if common else 0
        candidates.append({
            "signature": signature,
            "count": len(group),
            "avgMinutes": avg,
            "cadence": cadence,
            "commonSteps": [f"{s[0]} | {s[1]}" for s in common],
            "automationFit": fit,
            "score": round1(len(group) * avg * fit),
            "estMinutesSaved": js_round(avg * fit * 0.7),
            "evidence": [g["id"] for g in group[:3]],
        })
    candidates.sort(key=lambda c: (-c["score"], c["signature"]))
    return {
        "asOf": as_of,
        "windowDays": window_days,
        "itemsConsidered": considered,
        "candidates": candidates[:MAX_CANDIDATES],
        "coveredBySkills": covered_found,
    }


def read_rows(path, table_name):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        if table_name in ws.tables:
            cells = ws[ws.tables[table_name].ref]
            headers = [str(c.value) for c in cells[0]]
            return [
                {h: "" if c.value is None else str(c.value).strip() for h, c in zip(headers, row)}
                for row in cells[1:]
            ]
    sys.exit(f"table {table_name} not found in {path}")


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--as-of", required=True, help="YYYY-MM-DD")
    parser.add_argument("--window", type=int, default=30)
    args = parser.parse_args()
    result = mine(read_rows(args.workbook, "Items"), read_rows(args.workbook, "Skills"), args.as_of, args.window)
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
