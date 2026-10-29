#!/usr/bin/env python3
"""Check a ChiefOfStaff.xlsx against the schema and the agent's cell conventions.

Usage:
  python3 m365/scripts/lint_workbook.py path/to/ChiefOfStaff.xlsx

Download the live workbook from OneDrive and run this after an eval pass or
whenever the agent's writes look off. Exit code 1 means at least one ERROR.
"""

import re
import sys
from pathlib import Path

try:
    import yaml
    from openpyxl import load_workbook
except ImportError:
    sys.exit("PyYAML and openpyxl are required: pip install pyyaml openpyxl")

M365_DIR = Path(__file__).resolve().parent.parent
SCHEMA = M365_DIR / "schema" / "workbook.yaml"

DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
DATETIME_RE = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}(:\d{2})?$")
ITEM_ID_RE = re.compile(r"^I-\d{8}-\d{6}$")
FEEDBACK_ID_RE = re.compile(r"^F-\d{8}-\d{6}$")
LINK_RE = re.compile(r"^\[[a-z_]+\] .+ \| \S+$")
SIGNATURE_RE = re.compile(r"^[a-z0-9-]+\|[a-z0-9-]+\|[a-z0-9-]+$")
SLOT_RE = re.compile(r"^\d{2}:\d{2}-\d{2}:\d{2}$")
LINK_COLUMNS = ("SourceLink", "ContextLinks", "Artifacts")
CLOSED_STATUSES = {"done", "dropped"}
EXAMPLE_ID = "I-EXAMPLE"
CONFIDENTIAL_OUTCOME = "Completed — see linked artifacts"


class Report:
    def __init__(self):
        self.errors = []
        self.warnings = []

    def error(self, where, msg):
        self.errors.append(f"ERROR {where}: {msg}")

    def warn(self, where, msg):
        self.warnings.append(f"WARN  {where}: {msg}")


def text(value):
    return "" if value is None else str(value).strip()


def read_table(wb, table_name, spec, report):
    """Return (headers, rows) for an Excel table; rows are dicts keyed by header."""
    ws = wb[spec["sheet"]] if spec["sheet"] in wb.sheetnames else None
    if ws is None or table_name not in ws.tables:
        report.error(table_name, f"table not found on sheet '{spec['sheet']}'")
        return [], []
    cells = ws[ws.tables[table_name].ref]
    headers = [text(c.value) for c in cells[0]]
    rows = [dict(zip(headers, (c.value for c in row))) for row in cells[1:]]
    return headers, rows


def check_columns(table_name, spec, headers, rows, report):
    expected = [c["name"] for c in spec["columns"]]
    if headers != expected:
        missing = [h for h in expected if h not in headers]
        extra = [h for h in headers if h not in expected]
        report.error(table_name, f"headers differ from schema (missing {missing}, extra {extra})")

    key = spec["key"]
    seen = set()
    for i, row in enumerate(rows, start=2):
        key_value = text(row.get(key))
        where = f"{table_name}!{key_value or f'row {i}'}"
        if not key_value:
            if any(text(v) for v in row.values()):
                report.warn(where, f"row has data but no {key}; the agent will ignore it")
            continue
        if key_value in seen:
            report.error(where, f"duplicate {key}")
        seen.add(key_value)

        for col in spec["columns"]:
            value = text(row.get(col["name"]))
            if not value:
                continue
            if "values" in col and value not in col["values"]:
                report.error(where, f"{col['name']}='{value}' not in {col['values']}")
            kind = col.get("type")
            if kind == "int" and not re.fullmatch(r"-?\d+(\.0+)?", value):
                report.error(where, f"{col['name']}='{value}' is not a whole number")
            elif kind == "decimal" and not re.fullmatch(r"-?\d+(\.\d+)?", value):
                report.error(where, f"{col['name']}='{value}' is not a number")
            elif kind == "date" and not DATE_RE.match(value):
                report.error(where, f"{col['name']}='{value}' is not ISO YYYY-MM-DD text")
            elif kind == "datetime" and not DATETIME_RE.match(value):
                report.error(where, f"{col['name']}='{value}' is not ISO YYYY-MM-DDTHH:MM text")


def keyed(rows, key):
    return {text(r.get(key)): r for r in rows if text(r.get(key))}


def check_items(items, goal_ids, skill_names, report):
    open_count = linked_count = 0
    for item_id, row in keyed(items, "ID").items():
        where = f"Items!{item_id}"
        get = lambda col: text(row.get(col))
        if item_id != EXAMPLE_ID and not ITEM_ID_RE.match(item_id) and not item_id.startswith("I-SEED-"):
            report.warn(where, "ID doesn't match I-YYYYMMDD-HHMMSS")

        status, is_open = get("Status"), get("Open")
        if status and is_open:
            expected_open = "No" if status in CLOSED_STATUSES else "Yes"
            if is_open != expected_open:
                report.error(where, f"Open={is_open} but Status={status} (expected Open={expected_open})")
        if not get("Title"):
            report.error(where, "Title is empty")

        for col in LINK_COLUMNS:
            for line in filter(None, (l.strip() for l in get(col).splitlines())):
                if not LINK_RE.match(line):
                    report.error(where, f"{col} line not '[type] Title | url': {line[:60]}")
        if get("PlannedSlot") and not SLOT_RE.match(get("PlannedSlot")):
            report.error(where, f"PlannedSlot='{get('PlannedSlot')}' is not HH:MM-HH:MM")
        if get("Signature") and not SIGNATURE_RE.match(get("Signature")):
            report.error(where, f"Signature='{get('Signature')}' is not verb|object|counterpart")
        if get("Goal") and get("Goal") not in goal_ids and item_id != EXAMPLE_ID:
            report.error(where, f"Goal='{get('Goal')}' not found in Goals")
        if get("SkillUsed") and get("SkillUsed") not in skill_names:
            report.error(where, f"SkillUsed='{get('SkillUsed')}' not found in Skills")
        if get("Score") and get("Tier"):
            score = float(get("Score")) if re.fullmatch(r"-?\d+(\.\d+)?", get("Score")) else None
            if score is not None:
                tier = "P1" if score >= 70 else "P2" if score >= 50 else "P3" if score >= 30 else "P4"
                if tier != get("Tier"):
                    report.warn(where, f"Score {get('Score')} implies {tier}, but Tier={get('Tier')}")

        if status == "done":
            for col in ("CompletedOn", "Outcome", "Signature"):
                if not get(col):
                    report.error(where, f"done but {col} is empty")
            if not get("Steps"):
                report.warn(where, "done but Steps is empty (skill mining needs it)")

        if get("Confidential") == "Yes":
            for col in ("ResearchNotes", "Decisions"):
                if get(col):
                    report.error(where, f"Confidential=Yes but {col} has content")
            if get("Outcome") and get("Outcome") != CONFIDENTIAL_OUTCOME:
                report.error(where, f"Confidential=Yes but Outcome isn't '{CONFIDENTIAL_OUTCOME}'")

        if is_open == "Yes":
            open_count += 1
            if get("SourceLink") or get("ContextLinks"):
                linked_count += 1
            elif "no context found" not in get("ResearchNotes").lower():
                report.warn(where, "open item has no SourceLink/ContextLinks and no 'no context found' note")
    return open_count, linked_count


def check_skills(skills, report):
    for name, row in keyed(skills, "Name").items():
        where = f"Skills!{name}"
        if not re.fullmatch(r"[a-z0-9]+(-[a-z0-9]+)*", name):
            report.error(where, "Name must be kebab-case")
        runs, successes = text(row.get("Runs")) or "0", text(row.get("Successes")) or "0"
        if runs.isdigit() and successes.isdigit() and int(successes) > int(runs):
            report.error(where, f"Successes ({successes}) > Runs ({runs})")
        for sig in filter(None, (s.strip() for s in text(row.get("Signatures")).splitlines())):
            if not SIGNATURE_RE.match(sig):
                report.error(where, f"signature '{sig}' is not verb|object|counterpart")
        if text(row.get("Status")) == "active" and text(row.get("Tier")) == "playbook" and not text(row.get("Body")):
            report.error(where, "active playbook has no Body")


def check_feedback(feedback, report):
    for fid in keyed(feedback, "ID"):
        if not FEEDBACK_ID_RE.match(fid):
            report.warn(f"Feedback!{fid}", "ID doesn't match F-YYYYMMDD-HHMMSS")


def lint(path):
    schema = yaml.safe_load(SCHEMA.read_text(encoding="utf-8"))
    wb = load_workbook(path, data_only=True)
    report = Report()
    tables = {}
    for table_name, spec in schema["tables"].items():
        headers, rows = read_table(wb, table_name, spec, report)
        if headers:
            check_columns(table_name, spec, headers, rows, report)
        tables[table_name] = rows

    goal_ids = set(keyed(tables["Goals"], "ID"))
    skill_names = set(keyed(tables["Skills"], "Name"))
    open_count, linked_count = check_items(tables["Items"], goal_ids, skill_names, report)
    check_skills(tables["Skills"], report)
    check_feedback(tables["Feedback"], report)
    return report, open_count, linked_count


def main(argv):
    if len(argv) != 2:
        sys.exit(__doc__)
    report, open_count, linked_count = lint(argv[1])
    for line in report.errors + report.warnings:
        print(line)
    pct = f"{100 * linked_count // open_count}%" if open_count else "n/a"
    print(f"{len(report.errors)} errors, {len(report.warnings)} warnings · "
          f"open items: {open_count}, with ≥1 link: {pct}")
    return 1 if report.errors else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
