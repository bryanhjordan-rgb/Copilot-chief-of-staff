"""Tests for m365/scripts. Run: python3 -m unittest discover -s m365/tests -v"""

import importlib.util
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"


def load(name):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


skills = load("skills")
build_workbook = load("build_workbook")
lint_workbook = load("lint_workbook")


class WorkbookTestCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.path = self.tmp / "ChiefOfStaff.xlsx"
        build_workbook.build(self.path)

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def set_cells(self, sheet, row, values):
        """Write {header: value} into a data row of the table on `sheet`."""
        wb = load_workbook(self.path)
        ws = wb[sheet]
        headers = [c.value for c in ws[1]]
        for header, value in values.items():
            # .value assignment, not cell(value=...), so None clears the cell
            ws.cell(row=row, column=headers.index(header) + 1).value = value
        wb.save(self.path)

    def lint_errors(self):
        report, _, _ = lint_workbook.lint(self.path)
        return report.errors


class BuildWorkbookTests(WorkbookTestCase):
    def test_tables_and_keys_match_schema(self):
        wb = load_workbook(self.path)
        found = {name for ws in wb.worksheets for name in ws.tables}
        self.assertEqual(found, {"Items", "Goals", "Skills", "Feedback"})
        self.assertEqual(wb.sheetnames[0], "List")
        self.assertEqual(wb["List"]["A1"].value, "ID")

    def test_rebuild_is_byte_stable(self):
        second = self.tmp / "second.xlsx"
        build_workbook.build(second)
        self.assertEqual(self.path.read_bytes(), second.read_bytes())

    def test_dates_are_text_formatted(self):
        ws = load_workbook(self.path)["List"]
        headers = [c.value for c in ws[1]]
        due = ws.cell(row=2, column=headers.index("DueDate") + 1)
        self.assertEqual(due.number_format, "@")
        self.assertIsInstance(due.value, str)


class GoalsImportTests(WorkbookTestCase):
    GOALS = (
        "quarter: Q4 2026\nlast_updated: '2026-10-01'\n"
        "objectives:\n"
        "  - {name: Hire analysts, priority: 3, key_results: [Fill 2 roles], progress: 0.5, status: on_track}\n"
        "  - {name: Close renewals, priority: 1, target: $2M, key_results: [Contoso, Fabrikam], status: at_risk}\n"
        "personal:\n  - {name: Run 3x a week, progress: 0.2}\n"
    )

    def test_goals_are_ordered_and_personal_goals_prefixed(self):
        goals = self.tmp / "goals.yaml"
        goals.write_text(self.GOALS, encoding="utf-8")
        rows = build_workbook.goal_rows(goals)
        self.assertEqual([r["ID"] for r in rows], ["G1", "G2", "PG1"])
        self.assertEqual(rows[0]["Goal"], "Close renewals")
        self.assertEqual(rows[0]["KeyResults"], "Contoso\nFabrikam")
        self.assertEqual(rows[0]["Quarter"], "Q4 2026")

    def test_goals_land_in_table_and_lint_clean(self):
        goals = self.tmp / "goals.yaml"
        goals.write_text(self.GOALS, encoding="utf-8")
        build_workbook.build(self.path, goals)
        ws = load_workbook(self.path)["Goals"]
        self.assertEqual(ws.tables["Goals"].ref, "A1:I4")
        self.assertEqual(self.lint_errors(), [])


class LintWorkbookTests(WorkbookTestCase):
    def test_fresh_workbook_is_clean(self):
        self.assertEqual(self.lint_errors(), [])

    def test_open_must_match_status(self):
        self.set_cells("List", 2, {"Open": "Yes"})  # example row is Status=done
        self.assertTrue(any("Open=Yes but Status=done" in e for e in self.lint_errors()))

    def test_bad_link_line(self):
        self.set_cells("List", 2, {"ContextLinks": "see the email from Jane"})
        self.assertTrue(any("ContextLinks line" in e for e in self.lint_errors()))

    def test_non_iso_date(self):
        self.set_cells("List", 2, {"DueDate": "11/05/2026"})
        self.assertTrue(any("DueDate='11/05/2026'" in e for e in self.lint_errors()))

    def test_done_requires_signature(self):
        self.set_cells("List", 2, {"Signature": None})
        self.assertTrue(any("done but Signature is empty" in e for e in self.lint_errors()))

    def test_unknown_goal_on_real_item(self):
        self.set_cells("List", 2, {"ID": "I-20261029-081200", "Goal": "G9"})
        self.assertTrue(any("Goal='G9' not found" in e for e in self.lint_errors()))

    def test_confidential_hides_notes(self):
        self.set_cells("List", 2, {"Confidential": "Yes"})  # example has ResearchNotes
        errors = self.lint_errors()
        self.assertTrue(any("Confidential=Yes but ResearchNotes" in e for e in errors))

    def test_choice_value_outside_schema(self):
        self.set_cells("List", 2, {"Status": "finished"})
        self.assertTrue(any("Status='finished' not in" in e for e in self.lint_errors()))

    def test_skill_successes_cannot_exceed_runs(self):
        self.set_cells("Skills", 2, {"Name": "send-weekly-status", "Runs": 2, "Successes": 3})
        self.assertTrue(any("Successes (3) > Runs (2)" in e for e in self.lint_errors()))


class CardTests(unittest.TestCase):
    CARDS = Path(__file__).resolve().parent.parent / "flows" / "cards"
    ACTIONS = {
        "daily-brief": {"accept", "edit", "replan"},
        "weekly-review": {"decide", "approve_all", "reject_all"},
        "completion-confirm": {"complete", "not_yet"},
    }

    def strings(self, node):
        if isinstance(node, dict):
            for value in node.values():
                yield from self.strings(value)
        elif isinstance(node, list):
            for value in node:
                yield from self.strings(value)
        elif isinstance(node, str):
            yield node

    def test_cards(self):
        readme = (self.CARDS / "README.md").read_text(encoding="utf-8")
        for name, actions in self.ACTIONS.items():
            with self.subTest(card=name):
                card = json.loads((self.CARDS / f"{name}.card.json").read_text(encoding="utf-8"))
                self.assertEqual(card["type"], "AdaptiveCard")
                self.assertEqual(card["version"], "1.5")  # highest version Teams renders
                for s in self.strings(card):
                    if "{{" in s:
                        self.assertRegex(s, r"^\{\{[A-Z_]+\}\}$", "tokens must be whole string values")
                        self.assertIn(f"`{s[2:-2]}`", readme, "token must be documented")
                submitted = {a["data"]["action"] for a in card["actions"] if a["type"] == "Action.Submit"}
                self.assertEqual(submitted, actions)
                for action in actions:
                    self.assertIn(f"`{action}`", readme)


class SkillsValidatorTests(unittest.TestCase):
    def test_repo_skills_validate(self):
        self.assertTrue(skills.validate())

    def test_bad_skill_is_reported(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            bad = tmp / "Bad_Skill"
            bad.mkdir()
            (bad / "SKILL.md").write_text(
                "---\nname: Bad_Skill\ndescription: x\nmetadata:\n  version: 1\n"
                "  signatures: \"capture-item\"\n  tools: \"excel\"\n---\n# x\n## Goal\n## Steps\n",
                encoding="utf-8",
            )
            errors = " ".join(skills.validate_skill(bad))
            self.assertIn("kebab-case", errors)
            self.assertIn("metadata.version must be a quoted string", errors)
            self.assertIn("verb|object|counterpart", errors)
            self.assertIn("missing section '## Guardrails'", errors)
        finally:
            shutil.rmtree(tmp)

    def test_missing_evals_are_reported(self):
        self.assertEqual(skills.validate_evals("no-such-skill"), ["no-such-skill: missing evals/no-such-skill.yaml"])


if __name__ == "__main__":
    unittest.main()
